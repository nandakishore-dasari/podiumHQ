import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor

# ------------------------------------------------------------------------------
# Student-Level Slide Content Specification (10 Slides)
# ------------------------------------------------------------------------------

SLIDES_DATA = [
    {
        "num": 1,
        "title": "PodiumHQ — Hackathon Scoring Platform",
        "subtitle": "DevSecOps Capstone Project | Final Year Technical Presentation",
        "tag": "PROJECT INTRODUCTION",
        "cards": [
            {
                "header": "What is PodiumHQ?",
                "color": "#F59E0B",
                "items": [
                    "A full-stack web platform for organizing and judging hackathons.",
                    "Allows judges to evaluate teams across multiple rounds and weighted criteria.",
                    "Computes team rankings in real time and displays an automated live leaderboard."
                ]
            },
            {
                "header": "Why I Built This (Problem & Goal)",
                "color": "#10B981",
                "items": [
                    "Solves the mess of manual spreadsheets and delayed scores during hackathon deadlines.",
                    "Original application domain—built from scratch with zero TaskBoard code reuse.",
                    "Goal: Take a real web app through the complete DevSecOps pipeline from commit to cloud."
                ]
            },
            {
                "header": "Student & Project Info",
                "color": "#38BDF8",
                "items": [
                    "Student Name: [Your Name]",
                    "Course / Degree: B.Tech / B.E. Computer Science & Engineering",
                    "Stack: React 18, FastAPI, PostgreSQL, Docker, GitHub Actions, Terraform, Kubernetes"
                ]
            }
        ],
        "notes": (
            "Hi everyone. My name is [Your Name], and today I am presenting my DevSecOps capstone project, "
            "PodiumHQ. During college hackathons, I noticed organizers always struggle with messy spreadsheets "
            "and calculating weighted scores under tight deadlines. I built PodiumHQ to solve this: it is a web "
            "platform where judges submit scores against specific criteria and the leaderboard updates live. "
            "More importantly, I used this project to learn and implement a complete DevSecOps pipeline, including "
            "automated testing, non-root containers, CI/CD scanning, Terraform, Kubernetes, and monitoring."
        )
    },
    {
        "num": 2,
        "title": "Application Stack & Why I Chose These Tools",
        "subtitle": "Simple, clean 3-tier architecture built with modern open-source tools",
        "tag": "TECH STACK",
        "cards": [
            {
                "header": "Frontend: React 18 + Vite",
                "color": "#38BDF8",
                "items": [
                    "What it does: Single-page application with Tailwind CSS in an Obsidian & Solar Gold theme.",
                    "Why I chose it: Needed a fast, reactive UI where judges can submit scores without page refreshes.",
                    "Runtime: Served through unprivileged Nginx on port 8080 (running as UID 101)."
                ]
            },
            {
                "header": "Backend: FastAPI (Python 3.12)",
                "color": "#10B981",
                "items": [
                    "What it does: Handles 15+ REST endpoints for teams, criteria, judges, and score calculations.",
                    "Why I chose it: Automatic Swagger docs at /docs, async performance, and built-in health probes.",
                    "Runtime: Hardened Python container running strictly as non-root appuser (UID 10001)."
                ]
            },
            {
                "header": "Database: PostgreSQL 16",
                "color": "#F59E0B",
                "items": [
                    "What it does: Stores events, rounds, criteria, teams, judges, and evaluations across 6 tables.",
                    "Why I chose it: Strong relational integrity, foreign keys, and safe schema migrations with Alembic.",
                    "Connection: Managed with SQLAlchemy 2.0 ORM and Psycopg 3 connection pooling."
                ]
            }
        ],
        "notes": (
            "On Slide 2, here is how the application is built. For the frontend, I chose React 18 with Vite and "
            "Tailwind CSS so the scoring page and leaderboard update smoothly without refreshing. "
            "For the backend, I used FastAPI with Python 3.12 because it is asynchronous, validates incoming requests "
            "with Pydantic schemas, and automatically generates interactive API documentation at /docs. "
            "For storage, I used PostgreSQL 16 with SQLAlchemy and Alembic migrations to keep our data consistent. "
            "Both backend and frontend containers run as non-root users for basic security hygiene."
        )
    },
    {
        "num": 3,
        "title": "System Architecture: From Code to Deployment",
        "subtitle": "How code moves from local development to automated deployment and monitoring",
        "tag": "SYSTEM ARCHITECTURE",
        "diagram_boxes": [
            {"label": "1. Developer", "sub": "Git Push / Main", "color": "#64748B"},
            {"label": "2. GitHub Actions", "sub": "Lint -> Tests -> Scan", "color": "#F59E0B"},
            {"label": "3. Registry", "sub": "GHCR (Git SHA Tag)", "color": "#38BDF8"},
            {"label": "4. GitOps", "sub": "ArgoCD Auto-Sync", "color": "#10B981"},
            {"label": "5. AWS EKS", "sub": "VPC + Pods + HPA", "color": "#A855F7"},
            {"label": "6. Monitoring", "sub": "Prometheus & Grafana", "color": "#EC4899"}
        ],
        "cards": [
            {
                "header": "Delivery Flow (CI/CD to Registry)",
                "color": "#F59E0B",
                "items": [
                    "Developer pushes code to GitHub repository.",
                    "GitHub Actions automatically runs linters, 31 pytest tests, and builds container images.",
                    "Trivy scans images for security vulnerabilities before anything is published.",
                    "Images are tagged with the Git commit SHA (e.g. :f990516) and pushed to GHCR."
                ]
            },
            {
                "header": "Runtime Flow (Cluster & Monitoring)",
                "color": "#10B981",
                "items": [
                    "ArgoCD detects new image tags in Git and synchronizes manifests with the EKS cluster.",
                    "Ingress routes incoming traffic: / to frontend pods and /api to backend pods.",
                    "Horizontal Pod Autoscaler (HPA) scales backend pods between 2 and 6 replicas based on CPU.",
                    "Prometheus continuously scrapes /metrics to track response times and error rates."
                ]
            }
        ],
        "notes": (
            "This diagram on Slide 3 shows the overall architecture. "
            "When I push code to GitHub, GitHub Actions runs our automated pipeline: linting, running 31 pytest tests, "
            "building the Docker containers, and scanning them with Trivy. "
            "If everything passes, the image is tagged with the Git commit hash and pushed to GitHub Container Registry. "
            "From there, ArgoCD pulls the update into Kubernetes. An Ingress controller sends user traffic to our frontend, "
            "API calls to our backend pods, and Prometheus continuously monitors the whole system."
        )
    },
    {
        "num": 4,
        "title": "CI/CD Pipeline with GitHub Actions",
        "subtitle": "Automating code checks, test coverage, and security vulnerability scanning",
        "tag": "CI/CD WORKFLOW",
        "cards": [
            {
                "header": "Step 1: Code Quality & Linting",
                "color": "#38BDF8",
                "items": [
                    "What it does: Runs Ruff for Python code formatting and Mypy for static type validation.",
                    "Why it matters: Catches syntax errors, bad imports, and type mismatches before tests run.",
                    "Frontend: Runs npm run lint to catch unused variables and component bugs."
                ]
            },
            {
                "header": "Step 2: Automated Tests & Coverage Gate",
                "color": "#10B981",
                "items": [
                    "What it does: Executes 31 unit and integration tests using Pytest.",
                    "Coverage Gate: Pytest-cov enforces a strict 80% minimum coverage threshold.",
                    "Current Result: 90.99% overall code coverage. If coverage falls below 80%, the build fails."
                ]
            },
            {
                "header": "Step 3: Build & Security Scanning (Trivy)",
                "color": "#F59E0B",
                "items": [
                    "What it does: Builds multi-stage Docker images and scans them with Trivy.",
                    "Security Gate: Fails the pipeline if any HIGH or CRITICAL CVEs are detected.",
                    "Publishing: Successful builds are pushed to GHCR with immutable Git SHA tags, never :latest."
                ]
            }
        ],
        "notes": (
            "Slide 4 covers our CI/CD pipeline, configured in .github/workflows/ci-cd.yml. "
            "I divided it into three clear gates. First, Ruff and Mypy check code formatting and type safety. "
            "Second, Pytest runs all 31 tests and checks code coverage. We enforce an 80% threshold, and our tests "
            "currently hit 91%. If any test fails or coverage drops, the pipeline stops immediately. "
            "Third, we build the Docker images and run Trivy to scan for security vulnerabilities. "
            "Only clean images get tagged with the Git commit SHA and pushed to the container registry."
        )
    },
    {
        "num": 5,
        "title": "Infrastructure as Code: Terraform on AWS",
        "subtitle": "Writing code to provision cloud networks and Kubernetes instead of manual setup",
        "tag": "TERRAFORM IAC",
        "cards": [
            {
                "header": "Network: AWS VPC (vpc.tf)",
                "color": "#38BDF8",
                "items": [
                    "Spans 3 Availability Zones in AWS ap-south-1 (Mumbai) for reliability.",
                    "Creates 3 public subnets (for internet ingress) and 3 private subnets (for worker nodes).",
                    "Includes a NAT Gateway so private worker nodes can safely pull container images.",
                    "Tagged with kubernetes.io/role/elb so AWS can automatically manage load balancers."
                ]
            },
            {
                "header": "Compute: EKS Cluster (eks.tf & addons.tf)",
                "color": "#10B981",
                "items": [
                    "Deploys a managed Kubernetes v1.30 control plane with managed node groups.",
                    "Worker nodes: 2 to 4 t3.medium EC2 instances distributed across private subnets.",
                    "Add-on 1: AWS EBS CSI driver—required so PostgreSQL can attach persistent disk volumes.",
                    "Add-on 2: Metrics Server—required so the Horizontal Pod Autoscaler can read CPU usage."
                ]
            },
            {
                "header": "Honest Cloud Status & Teardown Runbook",
                "color": "#F59E0B",
                "items": [
                    "Status: Terraform files are written, formatted, and validated with terraform validate.",
                    "Cost Note: EKS and NAT Gateways cost ~$0.15/hour, so this was not left permanently running on AWS.",
                    "Safety: Provided a complete terraform destroy runbook in terraform/README.md for clean teardown."
                ]
            }
        ],
        "notes": (
            "On Slide 5, we look at our cloud infrastructure defined in the terraform/ directory. "
            "Instead of clicking around the AWS console, I wrote Terraform code to create an AWS VPC across 3 Availability "
            "Zones, with private subnets for security and public subnets for ingress. "
            "The EKS configuration provisions a Kubernetes 1.30 cluster with 2 to 4 worker nodes. "
            "I also added two essential EKS add-ons: the EBS CSI driver for database storage, and the Metrics Server for pod autoscaling. "
            "Because AWS EKS costs real money every hour, the code is fully validated, and I included a teardown runbook to destroy resources safely."
        )
    },
    {
        "num": 6,
        "title": "Kubernetes & Helm: Packaging Workloads",
        "subtitle": "Managing application deployments, services, scaling, and database migrations",
        "tag": "KUBERNETES & HELM",
        "cards": [
            {
                "header": "Workloads & High Availability",
                "color": "#38BDF8",
                "items": [
                    "Packaged as a Helm chart in helm/podiumhq with values-dev and values-prod.",
                    "Backend & Frontend Deployments: 2 replicas each with RollingUpdate strategy.",
                    "PodDisruptionBudget (PDB): Guarantees at least 1 backend pod stays online during node updates.",
                    "Horizontal Pod Autoscaler (HPA): Automatically scales backend from 2 to 6 pods when CPU > 60%."
                ]
            },
            {
                "header": "Container Security in Kubernetes",
                "color": "#10B981",
                "items": [
                    "Enforces securityContext on all pods: runAsNonRoot: true.",
                    "Backend runs as UID 10001; Frontend runs as Nginx UID 101.",
                    "Drops all Linux capabilities (capabilities: drop: [ALL]) to limit attack surface.",
                    "Ingress routes / to frontend, and /api, /health, /ready, /metrics to backend."
                ]
            },
            {
                "header": "Database Migration Job Hook",
                "color": "#F59E0B",
                "items": [
                    "The Problem: When 2+ backend pods boot at the same time, they race to run database migrations.",
                    "The Solution: Created a Helm pre-upgrade/pre-install Job hook (migration-job.yaml).",
                    "Alembic migrations run to completion first in an isolated job; application pods start only after it succeeds."
                ]
            }
        ],
        "notes": (
            "On Slide 6, we look at how the application is packaged for Kubernetes using Helm. "
            "Both backend and frontend are configured with 2 replicas each, protected by a PodDisruptionBudget and autoscaled by HPA. "
            "In Kubernetes, security is critical, so our pod manifests strictly enforce non-root execution and drop all Linux capabilities. "
            "One practical problem I solved was database migration race conditions: if two backend pods start at once, they clash "
            "over database schema locks. I resolved this by turning migrations into a Helm pre-install and pre-upgrade Job hook "
            "that finishes completely before any application pods boot."
        )
    },
    {
        "num": 7,
        "title": "GitOps with ArgoCD: Declarative Delivery",
        "subtitle": "Using Git as the single source of truth for cluster state and self-healing",
        "tag": "GITOPS CONTINUOUS DELIVERY",
        "cards": [
            {
                "header": "What is GitOps & Why ArgoCD?",
                "color": "#38BDF8",
                "items": [
                    "Concept: Instead of engineers running manual kubectl apply, Git holds the desired state.",
                    "ArgoCD controller runs inside the cluster and constantly compares live state against Git.",
                    "Configured in argocd/application.yaml watching the helm/podiumhq chart in this repo."
                ]
            },
            {
                "header": "Automated Sync & Pruning",
                "color": "#10B981",
                "items": [
                    "Automated Sync: When CI updates the image tag in Git, ArgoCD automatically applies the change.",
                    "Automated Pruning (prune: true): Removes old deleted resources from the cluster automatically.",
                    "Audit Trail: Every cluster change corresponds to a verifiable Git commit hash."
                ]
            },
            {
                "header": "Self-Healing Drift Demonstration",
                "color": "#F59E0B",
                "items": [
                    "What is drift?: When someone manually alters resources in the cluster using kubectl.",
                    "Self-Healing (selfHeal: true): ArgoCD detects the discrepancy and immediately reverts it.",
                    "Demo check: If I run 'kubectl scale deployment backend --replicas=5', ArgoCD forces it back to 2."
                ]
            }
        ],
        "notes": (
            "Slide 7 introduces GitOps using ArgoCD. "
            "Instead of manually deploying code with kubectl commands from a terminal, GitOps makes our Git repository "
            "the single source of truth. ArgoCD runs in the cluster and monitors our Helm chart. "
            "When our CI pipeline updates an image tag, ArgoCD automatically synchronizes the cluster. "
            "It also has self-healing enabled: if someone manually modifies a pod count or setting in the cluster, "
            "ArgoCD detects the drift and immediately restores the cluster back to what is defined in Git."
        )
    },
    {
        "num": 8,
        "title": "Observability: Prometheus & Grafana",
        "subtitle": "Monitoring system health, request traffic, and custom judging metrics",
        "tag": "OBSERVABILITY & METRICS",
        "cards": [
            {
                "header": "System Metrics on /metrics",
                "color": "#38BDF8",
                "items": [
                    "What it does: Backend exposes a Prometheus-format endpoint at /metrics.",
                    "HTTP Telemetry: Tracks total requests, response status codes (2xx, 4xx, 5xx), and p95 latency.",
                    "Runtime Telemetry: Tracks Python garbage collection, memory usage, CPU time, and open file descriptors."
                ]
            },
            {
                "header": "Custom Domain Business Metrics",
                "color": "#10B981",
                "items": [
                    "podiumhq_scores_submitted_total: Counter tracking how many judge scores have been submitted.",
                    "podiumhq_teams_by_status: Gauge showing active vs evaluated teams in the hackathon.",
                    "podiumhq_leaderboard_requests_total: Counter tracking how often the leaderboard is viewed."
                ]
            },
            {
                "header": "Grafana Dashboard & Alert Rules",
                "color": "#F59E0B",
                "items": [
                    "Grafana JSON (monitoring/grafana-dashboard.json): 5 panels showing RPS, p95 latency, 5xx errors, and scores.",
                    "PrometheusRule (monitoring/alerts.yaml): Triggers an alert if 5xx errors exceed 1% for 2 minutes.",
                    "Target Down Alert: Alerts immediately if the backend service stops responding."
                ]
            }
        ],
        "notes": (
            "On Slide 8, we look at observability. "
            "Under /metrics, our backend exposes standard Prometheus telemetry, including HTTP request rates and p95 latency histograms. "
            "To make monitoring genuinely useful for this hackathon platform, I also implemented three custom domain metrics: "
            "scores submitted, team status counts, and leaderboard views. "
            "We packaged a pre-built Grafana dashboard in the monitoring/ directory to visualize these numbers, "
            "along with Prometheus alert rules that warn us if error rates spike or if a service goes down."
        )
    },
    {
        "num": 9,
        "title": "Testing, Security & Real Debugging Challenges",
        "subtitle": "Verified test coverage, non-root checks, and authentic bugs I ran into and fixed",
        "tag": "TESTING & REAL DEBUGGING",
        "cards": [
            {
                "header": "Verified Test & Security Results",
                "color": "#10B981",
                "items": [
                    "Pytest Results: 31 tests passed in 1.32 seconds with 90.99% statement coverage.",
                    "Test Isolation: Runs against an in-memory SQLite database so tests never pollute real data.",
                    "Verified Non-Root: Ran 'docker compose exec backend id' -> uid=10001(appuser) and frontend -> uid=101(nginx)."
                ]
            },
            {
                "header": "Real Bug 1: In-Container Pytest Permission Crash",
                "color": "#F59E0B",
                "items": [
                    "The Problem: Running pytest inside the backend container crashed with sqlite3.OperationalError.",
                    "Root Cause: Working directory /app was created by Docker as root; appuser couldn't write the .coverage file.",
                    "The Fix: Added explicit 'RUN chown -R appuser:appgroup /app' in the Dockerfile."
                ]
            },
            {
                "header": "Real Bug 2: Nginx IPv6 Healthcheck Failure",
                "color": "#38BDF8",
                "items": [
                    "The Problem: Docker Compose reported the frontend container as 'unhealthy' even though web pages loaded.",
                    "Root Cause: BusyBox wget resolved localhost to IPv6 ::1, but Nginx was only listening on IPv4 8080.",
                    "The Fix: Added 'listen [::]:8080 default_server;' to the Nginx template and updated probe to 127.0.0.1."
                ]
            }
        ],
        "notes": (
            "Slide 9 highlights our testing and some real debugging problems I ran into and solved. "
            "Our test suite has 31 automated tests with 91% code coverage, and our containers were verified as non-root users. "
            "During development, I hit two interesting bugs. "
            "First, pytest kept crashing inside the container because the /app folder was owned by root, preventing our non-root user "
            "from writing the coverage database. I fixed this by chowning the directory in the Dockerfile. "
            "Second, Docker marked the frontend container unhealthy because BusyBox wget tried connecting to IPv6 localhost while "
            "Nginx only listened on IPv4. Adding an IPv6 listen directive solved the issue."
        )
    },
    {
        "num": 10,
        "title": "Summary, Limitations & Key Learnings",
        "subtitle": "Honest self-assessment of achievements, cloud limitations, and future improvements",
        "tag": "CONCLUSION & LEARNINGS",
        "cards": [
            {
                "header": "What Works & Was Verified",
                "color": "#10B981",
                "items": [
                    "Fully functional hackathon scoring platform running locally on Docker Compose.",
                    "31 automated tests passing with 90.99% coverage and isolated test fixtures.",
                    "Non-root container execution verified live on both backend and frontend.",
                    "Complete, validated manifest suite: Terraform, Helm chart, ArgoCD, and Prometheus."
                ]
            },
            {
                "header": "Honest Limitations (Cloud Scope)",
                "color": "#F59E0B",
                "items": [
                    "AWS EKS & ArgoCD were not deployed to live cloud infrastructure to avoid ongoing student cloud costs.",
                    "All cloud manifests are written, syntactically validated, and ready for deployment.",
                    "Everything runnable locally was verified 100% green with saved evidence files."
                ]
            },
            {
                "header": "Top Engineering Lessons",
                "color": "#38BDF8",
                "items": [
                    "DevSecOps is in the details: file permissions, non-root users, and dual-stack networking.",
                    "Helm pre-upgrade hooks are crucial to prevent database migration locks across replicas.",
                    "Immutable image tags (Git commit SHAs) prevent deployment caching confusion."
                ]
            }
        ],
        "notes": (
            "To conclude on Slide 10: PodiumHQ successfully demonstrates a modern DevSecOps project from scratch. "
            "Locally, everything is verified green: 31 tests pass with 91% coverage, all containers run healthy as non-root users, "
            "and all manifests for Terraform, Helm, ArgoCD, and Grafana are complete and validated. "
            "To be transparent about limitations: because cloud EKS and NAT Gateways incur hourly AWS bills, the cloud deployment "
            "was kept as validated code rather than a running cluster. "
            "This project gave me hands-on experience solving real configuration, security, and deployment challenges. "
            "Thank you, and I look forward to your questions!"
        )
    }
]

# ------------------------------------------------------------------------------
# 2. PowerPoint (.pptx) Builder
# ------------------------------------------------------------------------------

def build_pptx(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    COLOR_BG = RGBColor(8, 12, 20)          # #080C14 Obsidian
    COLOR_CARD = RGBColor(15, 23, 42)        # #0F172A Slate 900
    COLOR_BORDER = RGBColor(30, 41, 59)      # #1E293B Border
    COLOR_GOLD = RGBColor(245, 158, 11)      # #F59E0B Solar Gold
    COLOR_EMERALD = RGBColor(16, 185, 129)   # #10B981 Emerald
    COLOR_SKY = RGBColor(56, 189, 248)       # #38BDF8 Sky
    COLOR_TEXT = RGBColor(248, 250, 252)     # #F8FAFC Text light
    COLOR_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Muted text

    for slide_data in SLIDES_DATA:
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout

        # Add speaker notes
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = slide_data["notes"]

        # Background fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.color.rgb = COLOR_BG

        # Top Accent Bar
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.06))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_GOLD
        top_bar.line.color.rgb = COLOR_GOLD

        # Tag Badge
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(6.0), Inches(0.35))
        tf_tag = tag_box.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = f"STUDENT PRESENTATION • {slide_data['tag']}"
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_GOLD

        # Slide Number Badge (Right Aligned)
        num_box = slide.shapes.add_textbox(Inches(9.5), Inches(0.55), Inches(3.0), Inches(0.35))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"SLIDE {slide_data['num']} OF 10"
        p_num.font.name = "Arial"
        p_num.font.size = Pt(10)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_MUTED

        # Main Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = slide_data["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT

        # Subtitle Box
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = slide_data["subtitle"]
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = COLOR_EMERALD

        # Diagram Boxes (for Slide 3 architecture)
        if "diagram_boxes" in slide_data:
            box_w = Inches(1.8)
            box_h = Inches(0.85)
            spacing = Inches(0.18)
            start_x = Inches(0.8)
            start_y = Inches(2.0)

            for i, dbox in enumerate(slide_data["diagram_boxes"]):
                bx = start_x + i * (box_w + spacing)
                bshape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, start_y, box_w, box_h)
                bshape.fill.solid()
                bshape.fill.fore_color.rgb = COLOR_CARD
                bshape.line.color.rgb = COLOR_GOLD if i == 1 else COLOR_BORDER
                bshape.line.width = Pt(1.5) if i == 1 else Pt(1.0)

                tf_b = bshape.text_frame
                tf_b.word_wrap = True
                p1 = tf_b.paragraphs[0]
                p1.text = dbox["label"]
                p1.font.name = "Arial"
                p1.font.size = Pt(11)
                p1.font.bold = True
                p1.font.color.rgb = COLOR_TEXT
                p1.alignment = PP_ALIGN.CENTER

                p2 = tf_b.add_paragraph()
                p2.text = dbox["sub"]
                p2.font.name = "Calibri"
                p2.font.size = Pt(9)
                p2.font.color.rgb = COLOR_MUTED
                p2.alignment = PP_ALIGN.CENTER

            card_start_y = Inches(3.1)
            card_height = Inches(3.8)
        else:
            card_start_y = Inches(2.05)
            card_height = Inches(4.85)

        # Render Content Cards
        cards = slide_data["cards"]
        num_cards = len(cards)
        total_w = Inches(11.733)
        gap = Inches(0.25)
        card_w = (total_w - gap * (num_cards - 1)) / num_cards

        for c_idx, card in enumerate(cards):
            c_x = Inches(0.8) + c_idx * (card_w + gap)

            # Card Container
            card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, card_start_y, card_w, card_height)
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = COLOR_CARD
            card_shape.line.color.rgb = COLOR_BORDER
            card_shape.line.width = Pt(1.0)

            # Card Header
            header_box = slide.shapes.add_textbox(c_x + Inches(0.15), card_start_y + Inches(0.15), card_w - Inches(0.3), Inches(0.45))
            tf_h = header_box.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = card["header"]
            p_h.font.name = "Arial"
            p_h.font.size = Pt(14)
            p_h.font.bold = True
            hex_c = card["color"]
            if hex_c == "#F59E0B":
                p_h.font.color.rgb = COLOR_GOLD
            elif hex_c == "#10B981":
                p_h.font.color.rgb = COLOR_EMERALD
            elif hex_c == "#38BDF8":
                p_h.font.color.rgb = COLOR_SKY
            else:
                p_h.font.color.rgb = COLOR_TEXT

            # Bullet List
            body_box = slide.shapes.add_textbox(c_x + Inches(0.15), card_start_y + Inches(0.65), card_w - Inches(0.3), card_height - Inches(0.8))
            tf_b = body_box.text_frame
            tf_b.word_wrap = True

            for b_idx, item in enumerate(card["items"]):
                p_item = tf_b.add_paragraph() if b_idx > 0 else tf_b.paragraphs[0]
                p_item.text = "• " + item
                p_item.font.name = "Calibri"
                p_item.font.size = Pt(12)
                p_item.font.color.rgb = COLOR_TEXT
                p_item.space_after = Pt(8)

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        tf_f = footer_box.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = "PodiumHQ DevSecOps Capstone Presentation • Verified Codebase & Realistic Engineering"
        p_f.font.name = "Arial"
        p_f.font.size = Pt(9)
        p_f.font.color.rgb = COLOR_MUTED

    prs.save(output_path)
    print(f"[PPTX] Revised student presentation saved to: {output_path}")

# ------------------------------------------------------------------------------
# 3. PDF Generator matching identical 16:9 layout using ReportLab
# ------------------------------------------------------------------------------

def build_pdf(output_path):
    width = 13.333 * 72
    height = 7.5 * 72

    c = canvas.Canvas(output_path, pagesize=(width, height))

    C_BG = HexColor("#080C14")
    C_CARD = HexColor("#0F172A")
    C_BORDER = HexColor("#1E293B")
    C_GOLD = HexColor("#F59E0B")
    C_EMERALD = HexColor("#10B981")
    C_SKY = HexColor("#38BDF8")
    C_TEXT = HexColor("#F8FAFC")
    C_MUTED = HexColor("#94A3B8")

    for slide in SLIDES_DATA:
        # Background
        c.setFillColor(C_BG)
        c.rect(0, 0, width, height, fill=True, stroke=False)

        # Top Gold Bar
        c.setFillColor(C_GOLD)
        c.rect(0.8 * 72, height - 0.45 * 72, 11.733 * 72, 4, fill=True, stroke=False)

        # Category Tag
        c.setFillColor(C_GOLD)
        c.setFont("Helvetica-Bold", 10)
        c.drawString(0.8 * 72, height - 0.7 * 72, f"STUDENT PRESENTATION • {slide['tag']}")

        # Slide Number
        c.setFillColor(C_MUTED)
        c.drawRightString(width - 0.8 * 72, height - 0.7 * 72, f"SLIDE {slide['num']} OF 10")

        # Main Title
        c.setFillColor(C_TEXT)
        c.setFont("Helvetica-Bold", 24)
        c.drawString(0.8 * 72, height - 1.25 * 72, slide["title"])

        # Subtitle
        c.setFillColor(C_EMERALD)
        c.setFont("Helvetica", 13)
        c.drawString(0.8 * 72, height - 1.55 * 72, slide["subtitle"])

        # Handle Diagram Boxes on Slide 3
        if "diagram_boxes" in slide:
            box_w = 1.8 * 72
            box_h = 0.85 * 72
            gap_d = 0.18 * 72
            start_x = 0.8 * 72
            box_y = height - 2.65 * 72

            for i, dbox in enumerate(slide["diagram_boxes"]):
                bx = start_x + i * (box_w + gap_d)
                c.setFillColor(C_CARD)
                c.setStrokeColor(C_GOLD if i == 1 else C_BORDER)
                c.setLineWidth(1.5 if i == 1 else 1)
                c.roundRect(bx, box_y, box_w, box_h, 6, fill=True, stroke=True)

                c.setFillColor(C_TEXT)
                c.setFont("Helvetica-Bold", 10)
                c.drawCentredString(bx + box_w / 2.0, box_y + box_h - 22, dbox["label"])

                c.setFillColor(C_MUTED)
                c.setFont("Helvetica", 8)
                c.drawCentredString(bx + box_w / 2.0, box_y + 14, dbox["sub"])

            card_top = height - 3.0 * 72
            card_h = 3.7 * 72
        else:
            card_top = height - 1.95 * 72
            card_h = 4.8 * 72

        # Draw Cards
        cards = slide["cards"]
        num_cards = len(cards)
        total_w = 11.733 * 72
        gap_c = 0.25 * 72
        card_w = (total_w - gap_c * (num_cards - 1)) / num_cards

        for i, card in enumerate(cards):
            cx = 0.8 * 72 + i * (card_w + gap_c)
            cy = card_top - card_h

            c.setFillColor(C_CARD)
            c.setStrokeColor(C_BORDER)
            c.setLineWidth(1)
            c.roundRect(cx, cy, card_w, card_h, 8, fill=True, stroke=True)

            # Card Header
            h_color = card["color"]
            if h_color == "#F59E0B":
                c.setFillColor(C_GOLD)
            elif h_color == "#10B981":
                c.setFillColor(C_EMERALD)
            elif h_color == "#38BDF8":
                c.setFillColor(C_SKY)
            else:
                c.setFillColor(C_TEXT)

            c.setFont("Helvetica-Bold", 13)
            c.drawString(cx + 14, card_top - 28, card["header"])

            # Card Items
            c.setFillColor(C_TEXT)
            c.setFont("Helvetica", 9.5)
            text_y = card_top - 50

            for item in card["items"]:
                words = item.split(" ")
                lines = []
                curr_line = "• "
                for w in words:
                    if len(curr_line) + len(w) + 1 > 46:
                        lines.append(curr_line)
                        curr_line = "   " + w + " "
                    else:
                        curr_line += w + " "
                if curr_line:
                    lines.append(curr_line)

                for line in lines:
                    c.drawString(cx + 14, text_y, line)
                    text_y -= 14
                text_y -= 6

        # Footer
        c.setFillColor(C_MUTED)
        c.setFont("Helvetica", 8.5)
        c.drawString(0.8 * 72, 0.35 * 72, "PodiumHQ DevSecOps Capstone Presentation • Verified Codebase & Realistic Engineering")

        c.showPage()

    c.save()
    print(f"[PDF] Revised student presentation saved to: {output_path}")

# ------------------------------------------------------------------------------
# 4. Main Execution
# ------------------------------------------------------------------------------

if __name__ == "__main__":
    os.makedirs("demo", exist_ok=True)

    pptx_path = os.path.abspath("demo/presentation.pptx")
    pdf_path = os.path.abspath("demo/slides.pdf")
    alt_pptx_path = os.path.abspath("demo/slides.pptx")
    student_pptx_path = os.path.abspath("demo/presentation_student.pptx")

    # Build slides.pptx and presentation_student.pptx first (which are not locked)
    build_pptx(alt_pptx_path)
    build_pptx(student_pptx_path)

    # Try to update presentation.pptx if not open in WPS Office / PowerPoint
    try:
        build_pptx(pptx_path)
        print(f"[OK] Successfully updated: {pptx_path}")
    except PermissionError:
        print("[NOTICE] demo/presentation.pptx is currently open in WPS Office or locked.")
        print(f"         Revised student slides are saved at: {alt_pptx_path} and {student_pptx_path}")

    # Build PDF
    build_pdf(pdf_path)

    print("\nGeneration complete!")
    print(f"Presentation PPTX: {alt_pptx_path}")
    print(f"Presentation PDF:  {pdf_path}")
